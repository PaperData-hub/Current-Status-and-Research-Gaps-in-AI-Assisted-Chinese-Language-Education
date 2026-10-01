# -*- coding: utf-8 -*-
"""kappa_sample.py — 판정 일치도(원고 3.3, 부록 C)용 표본 추출과 코딩지 생성. 초록이 든 원자료가 필요하다.

1차 판정·코딩은 대규모 언어모델이 수행했으므로, 연구자가 AI 판정을 보지 않은 채 표본을 독립적으로 판정·코딩하고
두 결과의 Cohen's κ를 산출한다(kappa_compute.py).

표본(기본 200편, 난수 시드 20261001)
  - 포함 논문: 중국어군 전수(68편) + 영어군·한국어군 각 46편(연구방법 비율로 층화, 최대 나머지 배분)
  - 제외 논문: 40편(제외 사유 비율로 층화)
  - 행 순서는 무작위로 섞어 포함·제외나 언어군을 순서로 짐작할 수 없게 한다.
산출
  - data/kappa/κ표본_코딩지_<DATE>.xlsx   : 연구자가 기입할 파일(AI 판정값 없음)
  - data/kappa/κ표본_정답키_<DATE>.json   : 표본의 AI 판정값(코딩을 마칠 때까지 열어 보지 않는다)
사용: python scripts/rejudge/kappa_sample.py [--date 1161_20261001] [--n-en 46 --n-ko 46 --n-exc 40]
"""
import json, os, sys, io, random
from collections import Counter, defaultdict
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

REPO = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
arg = lambda k, d: type(d)(sys.argv[sys.argv.index(k) + 1]) if k in sys.argv else d
DATE = arg('--date', '1161_20261001')
N_EN, N_KO, N_EXC = arg('--n-en', 46), arg('--n-ko', 46), arg('--n-exc', 40)
SEED = 20261001
rng = random.Random(SEED)

A = json.load(open(f'{REPO}/data/rejudge/merged_{DATE}.json', encoding='utf-8'))
inc = [r for r in A if r['include'] == 'Y']; exc = [r for r in A if r['include'] != 'Y']


def strat(rows, key, n):
    """key별 비율로 n편을 나누고(최대 나머지 배분) 층 안에서 무작위 추출."""
    groups = defaultdict(list)
    for r in sorted(rows, key=lambda r: r['arti_id']):
        groups[key(r)].append(r)
    quota = {k: n * len(v) / len(rows) for k, v in groups.items()}
    alloc = {k: int(q) for k, q in quota.items()}
    for k in sorted(quota, key=lambda k: -(quota[k] - alloc[k]))[:n - sum(alloc.values())]:
        alloc[k] += 1
    out = []
    for k in sorted(groups):
        out += rng.sample(groups[k], alloc[k])
    return out


sample = [r for r in inc if r['group'] == '중국어']
sample += strat([r for r in inc if r['group'] == '영어'], lambda r: r['meth'], N_EN)
sample += strat([r for r in inc if r['group'] == '한국어'], lambda r: r['meth'], N_KO)
sample += strat(exc, lambda r: r['reason'], N_EXC)
rng.shuffle(sample)
print('표본', len(sample), '| 포함', Counter(r['group'] for r in sample if r['include'] == 'Y'), '| 제외', sum(r['include'] != 'Y' for r in sample))

os.makedirs(f'{REPO}/data/kappa', exist_ok=True)
KEYF = ['include', 'reason', 'group', 'learner', 'func', 'meth', 'tool']
key = [dict(no=i, arti_id=r['arti_id'], **{k: r[k] for k in KEYF}) for i, r in enumerate(sample, 1)]
json.dump(key, open(f'{REPO}/data/kappa/κ표본_정답키_{DATE}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
wb = Workbook(); ws0 = wb.active; ws0.title = '안내'
guide = [
    '코더 간 일치도(κ) 산출용 표본 코딩지',
    '',
    f'표본 {len(sample)}편. 각 논문의 제목·주제어·초록만 보고, 부록 A(판정 지침)의 기준에 따라 독립적으로 판정·코딩한다.',
    'AI의 1차 판정값은 이 파일에 없다. 판정을 마칠 때까지 data/kappa/κ표본_정답키 파일과 확정 코퍼스 엑셀을 열어 보지 않는다.',
    '',
    '기입 열(주황색). 칸을 누르면 허용값 목록이 나온다.',
    '  포함: Y(포함) / N(제외)',
    '  제외사유: N일 때만. 제2외국어, 언어비특정, 복수언어, 교육맥락없음, 번역품질, AI신호부재, 비AI에듀테크, 메타버스VR, 교양교육, 국어교과담론, AI디지털교과서, 기타',
    '  언어군: Y일 때만. 영어 / 한국어 / 중국어',
    '  학습자맥락: 한국어군일 때만. L1 / L2 / 미명시',
    '  언어기능(주)·연구방법(주)·핵심도구: Y일 때만. 부록 A의 <표 A3>·<표 A4>와 핵심 도구 목록',
    '  메모: 판정이 애매했던 이유 등(선택)',
    '',
    '기입을 마치면 파일을 저장하고 scripts/analysis/kappa_compute.py로 일치도를 계산한다.',
]
for t in guide:
    ws0.append([t])
ws0.column_dimensions['A'].width = 130; ws0['A1'].font = Font(bold=True, size=12)

ws = wb.create_sheet('코딩')
info = ['표본번호', 'artiId', '제목(국문)', '제목(영문)', '주제어(국문)', '주제어(영문)', '초록(국문)', '초록(영문)', '학술지', '연도']
code = ['포함', '제외사유', '언어군', '학습자맥락', '언어기능(주)', '연구방법(주)', '핵심도구', '메모']
ws.append(info + code)
HEAD = PatternFill('solid', fgColor='DDEBF7'); ORANGE = PatternFill('solid', fgColor='FCE4D6')
for c in range(1, len(info) + len(code) + 1):
    ws.cell(1, c).font = Font(bold=True); ws.cell(1, c).fill = ORANGE if c > len(info) else HEAD
for i, r in enumerate(sample, 1):
    ws.append([i, r['arti_id'], r['title_ko'], r['title_en'], r['kw_ko'], r['kw_en'], r['abstract_ko'], r['abstract_en'], r['journal'], r['year']] + [''] * len(code))
last = len(sample) + 1
LISTS = {
    '포함': 'Y,N',
    '제외사유': '제2외국어,언어비특정,복수언어,교육맥락없음,번역품질,AI신호부재,비AI에듀테크,메타버스VR,교양교육,국어교과담론,AI디지털교과서,기타',
    '언어군': '영어,한국어,중국어',
    '학습자맥락': 'L1,L2,미명시',
    '언어기능(주)': '쓰기,말하기,평가·채점,읽기,리터러시·역량,번역,문법,문학·문화,발음,어휘,듣기,기능비특정',
    '연구방법(주)': '개발연구,조사연구,질적·사례,실험연구,문헌·리뷰,성능평가,코퍼스분석,기타·혼합',
    '핵심도구': 'ChatGPT,생성형AI일반,중국어LLM,챗봇,기계번역,음성인식·TTS,BERT·NLP,기타',
}
for name, values in LISTS.items():
    col = get_column_letter(len(info) + code.index(name) + 1)
    dv = DataValidation(type='list', formula1=f'"{values}"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}{last}')
widths = {1: 8, 2: 15, 3: 40, 4: 30, 5: 25, 6: 25, 7: 70, 8: 50, 9: 18, 10: 7}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w
for c in range(len(info) + 1, len(info) + len(code) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 13
for row in ws.iter_rows(min_row=2, max_row=last, min_col=7, max_col=8):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'
out = f'{REPO}/data/kappa/κ표본_코딩지_{DATE}.xlsx'
wb.save(out)
print('saved', out)
