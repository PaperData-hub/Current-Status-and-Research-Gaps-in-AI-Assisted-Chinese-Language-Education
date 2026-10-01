# -*- coding: utf-8 -*-
"""kappa_compute.py — κ 표본의 확정 판정(정답키)과 독립 판정(연구자 코딩지 또는 AI 재판정 JSON)의 일치도를 산출(원고 3.3, 부록 C).
원고는 독립 재판정(다른 언어모델이 확정 판정값을 보지 않고 같은 지침으로 판정, 부록 C)과의 일치도를 보고한다:
  python scripts/analysis/kappa_compute.py --coded data/kappa/ai2/ai2_merged_1161_20261001.json --name1 "독립 재판정"
부록 C 본문은 appendix_c.py가 만든다.

축별 Cohen's κ = (관찰 일치도 − 우연 일치도) / (1 − 우연 일치도)
  - 포함 여부: 표본 전체
  - 언어군·언어기능·연구방법·핵심도구: 두 판정 모두 포함한 논문
  - 학습자 맥락: 두 판정 모두 한국어군으로 본 논문
  - 제외 사유(참고): 두 판정 모두 제외한 논문
두 번째 연구자의 코딩지(--coder2)를 주면 연구자 간 κ도 함께 산출한다.
산출: analysis/kappa_<DATE>.md(축별 κ, Landis–Koch 판정), analysis/kappa_불일치_<DATE>.md(부록 C용 불일치 목록)
사용: python scripts/analysis/kappa_compute.py [--date 1161_20261001] --coded <연구자1 코딩지> --coder2 <연구자2 코딩지>
     두 코딩지를 주면 AI–연구자1, AI–연구자2, 연구자1–연구자2의 세 쌍을 모두 산출한다.
"""
import json, os, sys, io
from collections import Counter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
REPO = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
arg = lambda k, d: sys.argv[sys.argv.index(k) + 1] if k in sys.argv else d
DATE = arg('--date', '1161_20261001')
CODED = arg('--coded', f'{REPO}/data/kappa/κ표본_코딩지_{DATE}_연구자1.xlsx')
CODER2 = arg('--coder2', None)
NAME1, NAME2 = arg('--name1', '연구자1'), arg('--name2', '연구자2')   # 예: --name1 "AI 2차 판정"
COLS = {'include': '포함', 'reason': '제외사유', 'group': '언어군', 'learner': '학습자맥락', 'func': '언어기능(주)', 'meth': '연구방법(주)', 'tool': '핵심도구'}


def read_coded(path):
    """코딩지(xlsx, '코딩' 시트) 또는 판정 결과 JSON 배열(id와 COLS의 키를 가진 객체)을 읽는다."""
    if path.lower().endswith('.json'):
        return {o['id']: {k: (str(o.get(k) or '').strip()) for k in COLS} | {'include': str(o.get('include') or '').strip().upper()}
                for o in json.load(open(path, encoding='utf-8'))}
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)['코딩']
    rows = list(ws.iter_rows(values_only=True)); h = list(rows[0])
    out = {}
    for r in rows[1:]:
        if r[h.index('artiId')] is None:
            continue
        d = {k: (str(r[h.index(v)]).strip() if r[h.index(v)] is not None else '') for k, v in COLS.items()}
        d['include'] = d['include'].upper()
        out[r[h.index('artiId')]] = d
    return out


def kappa(pairs):
    n = len(pairs)
    if not n:
        return None, None, 0
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n / n
    return ((po - pe) / (1 - pe) if pe < 1 else 1.0), po, n


def lk(k):
    if k is None: return '—'
    return '거의 완벽' if k > 0.80 else '상당함' if k > 0.60 else '보통' if k > 0.40 else '약함' if k > 0.20 else '미약'


def compare(A, B, la, lb):
    ids = [i for i in A if i in B]
    missing = [i for i in ids if not B[i]['include'] or not A[i]['include']]
    if missing:
        print(f'  ! 포함 여부가 빈 표본 {len(missing)}편(계산에서 제외):', missing[:5])
    ids = [i for i in ids if i not in missing]
    both_in = [i for i in ids if A[i]['include'] == 'Y' and B[i]['include'] == 'Y']
    both_ko = [i for i in both_in if A[i]['group'] == '한국어' and B[i]['group'] == '한국어']
    both_out = [i for i in ids if A[i]['include'] == 'N' and B[i]['include'] == 'N']
    axes = [('포함 여부', 'include', ids), ('언어군', 'group', both_in), ('언어기능(주)', 'func', both_in),
            ('연구방법(주)', 'meth', both_in), ('핵심도구', 'tool', both_in), ('학습자 맥락(한국어군)', 'learner', both_ko),
            ('제외 사유(참고)', 'reason', both_out)]
    res, dis = [], []
    for name, f, sub in axes:
        k, po, n = kappa([(A[i][f], B[i][f]) for i in sub])
        res.append((name, n, po, k))
        dis += [(name, i, A[i][f], B[i][f]) for i in sub if A[i][f] != B[i][f]]
    return res, dis


key = {o['arti_id']: {k: (o[k] or '') for k in COLS} for o in json.load(open(f'{REPO}/data/kappa/κ표본_정답키_{DATE}.json', encoding='utf-8'))}
title = {r['arti_id']: r['title_ko'] or r['title_en'] for r in json.load(open(f'{REPO}/data/rejudge/merged_{DATE}.json', encoding='utf-8')) if r['arti_id'] in key}
coders = [(NAME1, read_coded(CODED))] + ([(NAME2, read_coded(CODER2))] if CODER2 else [])
pairs = [('확정 판정', key, name, c) for name, c in coders]
if len(coders) == 2:
    pairs.append((NAME1, coders[0][1], NAME2, coders[1][1]))

L = [f'# 코더 간 일치도(κ) — 확정 코퍼스 {DATE}\n',
     f'표본 {len(key)}편. 비교 대상({NAME1}' + (f', {NAME2}' if CODER2 else '') + ')은 확정 판정값을 보지 않고 부록 A 지침에 따라 독립적으로 판정·코딩했다.\n']
D = [f'# 판정 불일치 사례 ({DATE})\n', '부록 C.3의 자료이다.\n']
for la, A_, lb, B_ in pairs:
    res, dis = compare(A_, B_, la, lb)
    L += [f'## {la} – {lb}\n', '| 축 | 대상 편수 | 관찰 일치도 | Cohen κ | Landis–Koch |', '|---|---|---|---|---|']
    for name, n, po, k in res:
        L.append(f'| {name} | {n} | {po * 100:.1f}% | {k:.2f} | {lk(k)} |' if k is not None else f'| {name} | 0 | — | — | — |')
    L.append('')
    D += [f'## {la} – {lb}: 불일치 {len(dis)}건\n', f'| 축 | artiId | {la} | {lb} | 제목 |', '|---|---|---|---|---|']
    D += [f'| {a} | {i} | {x or "(빈칸)"} | {y or "(빈칸)"} | {title.get(i, "")[:50]} |' for a, i, x, y in dis]
    D.append('')
io.open(f'{REPO}/analysis/kappa_{DATE}.md', 'w', encoding='utf-8', newline='\n').write('\n'.join(L) + '\n')
io.open(f'{REPO}/analysis/kappa_불일치_{DATE}.md', 'w', encoding='utf-8', newline='\n').write('\n'.join(D) + '\n')
print('\n'.join(L)); print(f'불일치 목록 → analysis/kappa_불일치_{DATE}.md')
